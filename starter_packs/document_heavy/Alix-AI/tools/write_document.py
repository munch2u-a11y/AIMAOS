import os

def _find_aimaos_root():
    p = os.path.dirname(os.path.abspath(__file__))
    while p != os.path.dirname(p) and not os.path.exists(os.path.join(p, "aimaos_config.yaml")):
        p = os.path.dirname(p)
    return p
AIMAOS_ROOT = os.environ.get("AIMAOS_ROOT") or _find_aimaos_root()
import sys
import shutil
import platform
import subprocess
import yaml
from datetime import datetime
from core.security import (
    SecurityValidationError,
    require_allowed_path,
    resolve_within,
    sanitize_output_basename,
    validate_slug,
)

# Optional imports handled gracefully
try:
    from docxtpl import DocxTemplate
except ImportError:
    DocxTemplate = None

TOOL_DEFINITION = {
    "name": "write_document",
    "description": "Fills a document template (.docx) with key-value fields and saves the output to the workspace output folder as a .docx or .pdf.",
    "parameters": {
        "type": "object",
        "properties": {
            "template_name": {
                "type": "string",
                "description": "The name of the template folder or template file (e.g. 'tax_return_1040' or 'tax_return_1040.docx')."
            },
            "field_values": {
                "type": "object",
                "description": "JSON object mapping placeholders to their replacement text (e.g. {'client_name': 'Alice'})."
            },
            "output_name": {
                "type": "string",
                "description": "Optional name for the output file (excluding extension). If not provided, it is auto-generated."
            },
            "output_format": {
                "type": "string",
                "enum": ["docx", "pdf"],
                "description": "The output file format. If not specified, defaults to config setting or docx."
            }
        },
        "required": ["template_name", "field_values"]
    }
}

def get_config():
    tools_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(tools_dir)
    config_path = os.path.join(project_dir, "config.yaml")
    if os.path.exists(config_path):
        with open(config_path, "r") as f:
            try:
                return yaml.safe_load(f)
            except Exception:
                pass
    return {}

def convert_docx_to_pdf(docx_path, output_dir):
    """Platform-agnostic DOCX to PDF converter."""
    # 1. Try Windows MS Word COM
    if platform.system() == "Windows":
        try:
            import docx2pdf
            pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
            docx2pdf.convert(docx_path, pdf_path)
            return pdf_path, "MS Word via docx2pdf"
        except Exception:
            pass

    # 2. Try LibreOffice / soffice
    libreoffice_bin = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice_bin:
        try:
            subprocess.run([
                libreoffice_bin,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                output_dir,
                docx_path
            ], check=True, capture_output=True)
            pdf_name = os.path.splitext(os.path.basename(docx_path))[0] + ".pdf"
            return os.path.join(output_dir, pdf_name), "LibreOffice headless"
        except Exception:
            pass

    # 3. Try ReportLab fallback (creates a simple plain text PDF)
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        import docx
        
        pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
        doc = docx.Document(docx_path)
        c = canvas.Canvas(pdf_path, pagesize=letter)
        width, height = letter
        
        y = height - 50
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, y, f"Document: {os.path.basename(docx_path)}")
        y -= 30
        c.setFont("Helvetica", 10)
        
        for para in doc.paragraphs:
            if para.text.strip():
                text = para.text
                # Wrap text helper
                words = text.split()
                line = []
                for word in words:
                    line.append(word)
                    # Check length
                    line_str = " ".join(line)
                    if c.stringWidth(line_str, "Helvetica", 10) > (width - 100):
                        # Draw previous line
                        c.drawString(50, y, " ".join(line[:-1]))
                        y -= 15
                        if y < 50:
                            c.showPage()
                            y = height - 50
                        line = [word]
                if line:
                    c.drawString(50, y, " ".join(line))
                    y -= 20
                    if y < 50:
                        c.showPage()
                        y = height - 50
                        
        c.save()
        return pdf_path, "ReportLab fallback"
    except Exception:
        pass
        
    return None, None

def execute(template_name, field_values, output_name=None, output_format=None):
    if not DocxTemplate:
        return "Error: docxtpl is not installed. Run 'pip install docxtpl'."

    config = get_config() or {}
    paths = config.get("paths", {})
    templates_dir = require_allowed_path(paths.get("templates", "./templates"))
    output_dir = require_allowed_path(paths.get("output", "./workspace/output"), must_exist=False)
    if not isinstance(field_values, dict) or len(field_values) > 100:
        return "Error: field_values must be an object with at most 100 fields."
    
    if not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)

    # Resolve an identifier, never an arbitrary model-provided path.
    template_id = os.path.splitext(os.path.basename(str(template_name)))[0]
    try:
        template_id = validate_slug(template_id, label="template identifier")
    except SecurityValidationError as exc:
        return f"Error: {exc}"
    try_dir = resolve_within(templates_dir, template_id, "template.docx")
    try_file = resolve_within(templates_dir, f"{template_id}.docx")
        
    if os.path.exists(try_dir):
        docx_template_path = try_dir
    elif os.path.exists(try_file):
        docx_template_path = try_file
    else:
        return f"Error: Template '{template_name}' not found. Searched:\n- {try_dir}\n- {try_file}"

    # Determine output filename
    if not output_name:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_name = f"{template_id}_{timestamp}"

    # Ensure format selection
    if not output_format:
        output_format = config.get("default_output_format", "docx")
    if output_format not in {"docx", "pdf"}:
        return "Error: output_format must be 'docx' or 'pdf'."
    output_name = sanitize_output_basename(str(output_name), fallback=f"{template_id}_draft")
    output_docx_path = resolve_within(output_dir, f"{output_name}.docx")

    # Render template using docxtpl (Jinja2 syntax)
    try:
        doc = DocxTemplate(docx_template_path)
        doc.render(field_values)
        doc.save(output_docx_path)
    except Exception as e:
        return f"Error rendering template: {e}"

    # Auto-log to DocumentProductionMemory
    try:
        sys.path.insert(0, os.path.join(AIMAOS_ROOT, "Alix-AI"))
        from business.memory import DocumentProductionMemory
        mem_dir = paths.get("memory", "./workspace/.memory")
        mem = DocumentProductionMemory(memory_dir=mem_dir)
        final_file = output_docx_path
        if output_format == "pdf":
            pdf_p, _ = convert_docx_to_pdf(output_docx_path, output_dir)
            if pdf_p:
                final_file = pdf_p
        mem.log_production(
            input_file="intake_document",
            template_used=template_name,
            output_file=final_file,
            extracted_fields=field_values,
            status="success"
        )
    except Exception as e:
        print(f"Warning: Could not log production to memory: {e}")

    if output_format == "docx":
        return f"Success: Template filled and saved to {output_docx_path}"
    
    elif output_format == "pdf":
        pdf_path, converter_used = convert_docx_to_pdf(output_docx_path, output_dir)
        if pdf_path:
            # We successfully converted to PDF! Optionally keep or remove docx
            # Let's keep it but tell the user about both.
            return f"Success: Template filled and converted to PDF using {converter_used}.\n- Word file: {output_docx_path}\n- PDF file: {pdf_path}"
        else:
            return f"Warning: Word file saved to {output_docx_path}, but PDF conversion failed (no Word processor or ReportLab PDF library detected on host)."
            
    return f"Error: Unsupported output format '{output_format}'."
