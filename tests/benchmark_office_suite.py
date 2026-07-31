"""
AIMAOS Office Suite Benchmark Harness
=====================================
Benchmarks B1-B7 derived from the AIMAOS Flaw Report
(System Technical Documents/AIMAOS_flaw_report_and_benchmarks.md).

Run with the Alix venv:
    python3 tests/benchmark_office_suite.py

Writes raw results to tests/benchmark_results.json.
"""
import os

def _find_aimaos_root():
    p = os.path.dirname(os.path.abspath(__file__))
    while p != os.path.dirname(p) and not os.path.exists(os.path.join(p, "aimaos_config.yaml")):
        p = os.path.dirname(p)
    return p
AIMAOS_ROOT = os.environ.get("AIMAOS_ROOT") or _find_aimaos_root()
import re
import sys
import json
import time
import glob
import shutil
import importlib.util
import multiprocessing
from datetime import datetime

AIMAOS = AIMAOS_ROOT
sys.path.insert(0, AIMAOS)

from core.comms.office_board import OfficeBoard, OFFICE_BOARD_FILE
from core.comms.bus import AgentCompanyBus

RESULTS = {"run_at": datetime.now().isoformat(), "benchmarks": {}}


def load_module(mod_name, file_path):
    spec = importlib.util.spec_from_file_location(mod_name, file_path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[mod_name] = mod
    spec.loader.exec_module(mod)
    return mod


def snapshot_board():
    """Backs up the live office board so benchmarks run on a clean board."""
    backup = OFFICE_BOARD_FILE + ".bench_backup"
    if os.path.exists(OFFICE_BOARD_FILE):
        shutil.copy(OFFICE_BOARD_FILE, backup)
        os.remove(OFFICE_BOARD_FILE)
    return backup


def restore_board(backup):
    if os.path.exists(backup):
        shutil.move(backup, OFFICE_BOARD_FILE)


# ---------------------------------------------------------------- B1
def bench_task_lifecycle():
    """12 mixed-priority tasks -> Marley dispatch -> agent turns.
    Measures completion rate, dispatch priority order, re-dispatch, stranding."""
    print("\n=== B1: TASK LIFECYCLE PIPELINE ===")
    board = OfficeBoard()
    orch_mod = load_module("bench_orch", f"{AIMAOS}/Marley-AI/core/orchestrator.py")
    alix_mod = load_module("bench_alix", f"{AIMAOS}/Alix-AI/core/agent.py")

    priorities = ["BACKGROUND", "HIGH", "NORMAL", "CRITICAL"] * 3
    posted = []
    for i, prio in enumerate(priorities):
        tid = board.post_task(f"Bench task {i} ({prio})", "Benchmark", "Alix", prio)
        posted.append((tid, prio))

    orchestrator = orch_mod.MarleyOrchestrator()
    dispatch_order = []
    seen_dispatch = set()
    redispatch_count = 0

    alix = alix_mod.AlixAgent()
    for _ in range(len(posted) + 4):  # extra cycles to detect re-dispatch
        turn = orchestrator.dispatch_next_turn()
        if isinstance(turn, dict):
            tid = turn["task_id"]
            if tid in seen_dispatch:
                redispatch_count += 1
            seen_dispatch.add(tid)
            dispatch_order.append(turn["priority"])
            alix.execute_single_turn()  # worker picks up and completes

    # Priority-order correctness: dispatches must be non-decreasing in weight
    weights = {"CRITICAL": 0, "HIGH": 1, "NORMAL": 2, "BACKGROUND": 3}
    order_ok = all(weights[dispatch_order[i]] <= weights[dispatch_order[i + 1]]
                   for i in range(len(dispatch_order) - 1))

    board.board = board._load_board()
    remaining = [t for t in board.board["active_tasks"] if t["requester"] == "Benchmark"]
    completed = [t for t in board.board["completed_tasks"] if t.get("requester") == "Benchmark"]

    res = {
        "tasks_posted": len(posted),
        "tasks_completed": len(completed),
        "completion_rate": round(len(completed) / len(posted), 3),
        "dispatch_priority_order_correct": order_ok,
        "dispatch_sequence": dispatch_order,
        "redispatch_count": redispatch_count,
        "stranded_tasks": len(remaining),
    }
    print(json.dumps(res, indent=2))
    return res


# ---------------------------------------------------------------- B2
def _writer_proc(worker_id, n):
    sys.path.insert(0, AIMAOS_ROOT)
    from core.comms.office_board import OfficeBoard
    board = OfficeBoard()
    for i in range(n):
        board.post_task(f"Concurrent task w{worker_id}-{i}", f"Worker{worker_id}", "Kai", "NORMAL")
        board.log_activity(f"Worker{worker_id} activity {i}")


def bench_concurrent_writers(workers=4, per_worker=15):
    """4 processes hammer the board simultaneously; count lost updates."""
    print("\n=== B2: CONCURRENT BOARD WRITERS ===")
    procs = [multiprocessing.Process(target=_writer_proc, args=(w, per_worker))
             for w in range(workers)]
    t0 = time.time()
    for p in procs:
        p.start()
    for p in procs:
        p.join()
    elapsed = time.time() - t0

    board = OfficeBoard()
    present = [t for t in board.board["active_tasks"] if t["title"].startswith("Concurrent task")]
    expected = workers * per_worker
    res = {
        "writer_processes": workers,
        "tasks_expected": expected,
        "tasks_present": len(present),
        "lost_updates": expected - len(present),
        "elapsed_sec": round(elapsed, 2),
    }
    # cleanup benchmark tasks
    for t in present:
        board.update_task_status(t["id"], "completed", result="bench cleanup")
    print(json.dumps(res, indent=2))
    return res


# ---------------------------------------------------------------- B3
def bench_specialist_honesty():
    """Distinct topics -> Quinn. If outputs are canned, differentiation ~ 0."""
    print("\n=== B3: SPECIALIST HONESTY (Quinn & Zoe) ===")
    quinn_tool = load_module("bench_quinn_brief", f"{AIMAOS}/Quinn-AI/tools/research_brief.py")
    topics = [
        "Florida Statute 61.13 child custody time-sharing factors",
        "Chapter 744 guardianship of incapacitated adults",
        "Florida Statute 732.502 execution requirements for wills",
    ]
    bodies = []
    for t in topics:
        quinn_tool.execute(t, scope="statutory")
        reports = sorted(glob.glob(f"{AIMAOS}/Quinn-AI/workspace/reports/brief_*.md"))
        with open(reports[-1]) as f:
            text = f.read()
        # Strip the echoed topic/timestamp header; compare the substantive body
        body = "\n".join(l for l in text.splitlines()
                         if t not in l and "Generated" not in l)
        bodies.append(body)
        time.sleep(1.1)  # distinct filenames (second-resolution)

    identical_pairs = sum(1 for i in range(len(bodies)) for j in range(i + 1, len(bodies))
                          if bodies[i] == bodies[j])
    total_pairs = len(bodies) * (len(bodies) - 1) // 2
    differentiation = 1 - (identical_pairs / total_pairs)

    # Zoe: does the improvement report reference any actual trace content?
    zoe_mod = load_module("bench_zoe_synth", f"{AIMAOS}/Zoe-AI/core/workflow_synthesizer.py")
    zoe_mod.ZoeWorkflowSynthesizer().synthesize_improvement_report()
    zreports = sorted(glob.glob(f"{AIMAOS}/Zoe-AI/workspace/diagnostics/*.md"))
    with open(zreports[-1]) as f:
        zoe_text = f.read()
    logs_dir = f"{AIMAOS}/comms/task_logs"
    trace_titles = []
    for lf in glob.glob(f"{logs_dir}/*.json"):
        try:
            with open(lf) as f:
                trace_titles.append(json.load(f).get("title") or "")
        except Exception:
            pass
    zoe_references_traces = any(t and t in zoe_text for t in trace_titles)
    zoe_claims_perfect = "100%" in zoe_text

    res = {
        "quinn_topics_tested": len(topics),
        "quinn_identical_body_pairs": f"{identical_pairs}/{total_pairs}",
        "quinn_differentiation_ratio": round(differentiation, 3),
        "zoe_report_references_actual_traces": zoe_references_traces,
        "zoe_hardcodes_perfect_efficiency": zoe_claims_perfect,
    }
    print(json.dumps(res, indent=2))
    return res


# ---------------------------------------------------------------- B4
def bench_ipc_roundtrip(messages_per_agent=10):
    """Fan 30 messages to Kai/Quinn/Zoe, process, measure replies + hygiene."""
    print("\n=== B4: IPC ROUND-TRIP ===")
    agents = {}
    for name in ["Kai", "Quinn", "Zoe"]:
        mod = load_module(f"bench_{name.lower()}_agent", f"{AIMAOS}/{name}-AI/core/agent.py")
        agents[name] = getattr(mod, f"{name}Agent")()

    sender = AgentCompanyBus("Alix")
    # Drain any stale replies first
    sender.read_inbox(mark_read=True)

    t0 = time.time()
    sent = 0
    for name in agents:
        for i in range(messages_per_agent):
            sender.send_message(name, "bench_ping", {"seq": i})
            sent += 1

    # Poison message: invalid JSON dropped into Kai's inbox
    poison_path = f"{AIMAOS}/comms/Kai/inbox/zz_poison_bench.json"
    with open(poison_path, "w") as f:
        f.write("{not valid json")

    for agent in agents.values():
        agent.process_inter_agent_messages()

    replies = [m for m in sender.read_inbox(mark_read=True)
               if m.get("action") == "reply_bench_ping"]
    elapsed = time.time() - t0

    poison_still_pending = os.path.exists(poison_path)
    if poison_still_pending:
        os.remove(poison_path)

    read_files = len(glob.glob(f"{AIMAOS}/comms/*/inbox/*.read"))
    res = {
        "messages_sent": sent,
        "replies_received": len(replies),
        "reply_rate": round(len(replies) / sent, 3),
        "elapsed_sec": round(elapsed, 3),
        "mean_roundtrip_ms": round(1000 * elapsed / sent, 1),
        "poison_message_requeued_forever": poison_still_pending,
        "accumulated_read_files_systemwide": read_files,
    }
    print(json.dumps(res, indent=2))
    return res


# ---------------------------------------------------------------- B5
def bench_security_triage():
    """Spoofed/legit senders -> Finn triage; measure classification accuracy."""
    print("\n=== B5: SECURITY TRIAGE ===")
    triage = load_module("bench_triage", f"{AIMAOS}/Finn-AI/tools/triage_incoming.py")
    cases = [
        ("client@gmail.com", True),
        ("clerk@court.fl.gov", True),
        ("partner@sub.lawfirm.com", True),          # legitimate subdomain
        ("evil@gmail.com.attacker.net", False),      # suffix spoof
        ("gmail.com@phishing.io", False),            # local-part spoof
        ("attacker@notgmail.com", False),            # lookalike domain
        ("randomperson@yahoo.com", False),           # simply not allowlisted
        ("spoof@courtxfl.gov", False),               # near-miss domain
    ]
    correct = 0
    details = []
    board = OfficeBoard()
    for addr, expected_verified in cases:
        out = triage.execute(sender_address=addr, message="benchmark probe", channel="email")
        got_verified = "Security Status: VERIFIED" in out
        ok = got_verified == expected_verified
        correct += ok
        details.append({"sender": addr, "expected": expected_verified,
                        "got": got_verified, "correct": ok})
    # cleanup probe tasks
    board.board = board._load_board()
    for t in list(board.board["active_tasks"]):
        if t.get("details", {}).get("message") == "benchmark probe":
            board.update_task_status(t["id"], "completed", result="bench cleanup")

    res = {
        "cases": len(cases),
        "correct": correct,
        "accuracy": round(correct / len(cases), 3),
        "details": details,
    }
    print(json.dumps(res, indent=2))
    return res


# ---------------------------------------------------------------- B6
def bench_document_production():
    """Render real court templates; check unresolved Jinja2 placeholders."""
    print("\n=== B6: DOCUMENT PRODUCTION ===")
    sys.path.insert(0, os.path.join(AIMAOS_ROOT, "Alix-AI"))
    from business.document_engine import DocumentEngine
    import docx as docx_lib

    clients = [
        ("Benchmark Client One", "form_12_982_a",
         {"client_name": "Benchmark Client One", "county": "Leon", "circuit_number": "2nd",
          "case_number": "2026-DR-1111", "new_name": "Benchmark Sterling",
          "client_address": "1 Test Way", "client_phone": "(850) 555-0000",
          "client_email": "bench@example.com", "date_of_birth": "Jan 1, 1990"}),
        ("Benchmark Client Two", "form_12_982_b",
         {"client_name": "Benchmark Client Two", "county": "Duval", "circuit_number": "4th",
          "case_number": "2026-DR-2222", "new_name": "Bench Mark II"}),
        ("Benchmark Client Three", "form_12_902_e",
         {"client_name": "Benchmark Client Three", "county": "Clay", "circuit_number": "4th",
          "case_number": "2026-DR-3333"}),
    ]
    out_root = f"{AIMAOS}/Alix-AI/workspace/output/_benchmark"
    os.makedirs(out_root, exist_ok=True)

    rendered, failures, unresolved_total = 0, [], 0
    t0 = time.time()
    for cname, form, ctx in clients:
        tpl = f"{AIMAOS}/Alix-AI/templates/{form}/template.docx"
        if not os.path.exists(tpl):
            failures.append(f"{form}: template missing")
            continue
        out = os.path.join(out_root, f"{form}_{cname.split()[-1].lower()}.docx")
        try:
            DocumentEngine(tpl).generate(ctx, out)
            rendered += 1
            d = docx_lib.Document(out)
            text = "\n".join(p.text for p in d.paragraphs)
            unresolved = re.findall(r"\{\{[^}]*\}\}", text)
            unresolved_total += len(unresolved)
        except Exception as e:
            failures.append(f"{form}: {e}")
    elapsed = time.time() - t0

    res = {
        "templates_attempted": len(clients),
        "rendered_ok": rendered,
        "render_failures": failures,
        "unresolved_placeholders_in_output": unresolved_total,
        "elapsed_sec": round(elapsed, 2),
    }
    print(json.dumps(res, indent=2))
    return res


# ---------------------------------------------------------------- B7
def bench_real_llm_turn():
    """The feasibility proof: office minimal prompt + tool schema against a
    real installed local model via the office's own LLMClient."""
    print("\n=== B7: REAL SINGLE-THOUGHT LLM TURN (qwen3.5:2b) ===")
    from core.llm import LLMClient
    from core.mrag_beliefs import AgentBeliefStore

    config = {"llm": {"backend": "ollama", "model": "qwen3.5:2b",
                      "temperature": 0.1, "max_tokens": 512}}
    client = LLMClient(config)
    ok, msg = client.check_availability()
    if not ok:
        res = {"available": False, "detail": msg}
        print(json.dumps(res, indent=2))
        return res

    store = AgentBeliefStore()
    sys_prompt = store.get_minimal_system_prompt("Alix", "Document Production & Keeper Agent")

    tool_def = load_module("bench_pop_tpl", f"{AIMAOS}/Alix-AI/tools/populate_template.py").TOOL_DEFINITION
    user_task = ("Office Board task: Generate the Adult Name Change Petition for client "
                 "Sam Doe in Leon County (2nd circuit), new name Sam Newname, "
                 "case number 2026-DR-4567, using template form_12_982_a. "
                 "Call the populate_template tool with the correct arguments.")

    t0 = time.time()
    try:
        resp = client.chat(
            [{"role": "system", "content": sys_prompt},
             {"role": "user", "content": user_task}],
            tools=[tool_def])
        elapsed = time.time() - t0
    except Exception as e:
        res = {"available": True, "turn_succeeded": False, "error": str(e)}
        print(json.dumps(res, indent=2))
        return res

    made_tool_call = bool(resp.tool_calls)
    correct_tool = made_tool_call and resp.tool_calls[0]["name"] == "populate_template"
    args = resp.tool_calls[0]["arguments"] if made_tool_call else {}
    ctx = args.get("context", {}) if isinstance(args.get("context"), dict) else {}
    correct_template = args.get("template_name") == "form_12_982_a"
    client_name_ok = "Sam Doe" in json.dumps(args)

    res = {
        "available": True,
        "model": client.model,
        "system_prompt_chars": len(sys_prompt),
        "turn_succeeded": True,
        "latency_sec": round(elapsed, 2),
        "made_tool_call": made_tool_call,
        "called_correct_tool": correct_tool,
        "template_name_correct": correct_template,
        "client_name_present_in_args": client_name_ok,
        "raw_arguments": args,
        "content_preview": (resp.content or "")[:200],
    }
    print(json.dumps(res, indent=2))
    return res


# ---------------------------------------------------------------- main
def main():
    print("====================================================================")
    print("AIMAOS OFFICE SUITE BENCHMARKS (B1-B7)")
    print("====================================================================")
    backup = snapshot_board()
    try:
        RESULTS["benchmarks"]["B1_task_lifecycle"] = bench_task_lifecycle()
        RESULTS["benchmarks"]["B2_concurrent_writers"] = bench_concurrent_writers()
        RESULTS["benchmarks"]["B3_specialist_honesty"] = bench_specialist_honesty()
        RESULTS["benchmarks"]["B4_ipc_roundtrip"] = bench_ipc_roundtrip()
        RESULTS["benchmarks"]["B5_security_triage"] = bench_security_triage()
        RESULTS["benchmarks"]["B6_document_production"] = bench_document_production()
        RESULTS["benchmarks"]["B7_real_llm_turn"] = bench_real_llm_turn()
    finally:
        restore_board(backup)

    out_path = f"{AIMAOS}/tests/benchmark_results.json"
    with open(out_path, "w") as f:
        json.dump(RESULTS, f, indent=2)
    print("\n====================================================================")
    print(f"BENCHMARKS COMPLETE. Raw results: {out_path}")
    print("====================================================================")


if __name__ == "__main__":
    main()
